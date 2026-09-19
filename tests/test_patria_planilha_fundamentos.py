from __future__ import annotations

import dataclasses
import datetime

import openpyxl
import pytest

from iip.cli import fetch_template
from iip.sources import patria_planilha_fundamentos_harvester as harvester_module
from iip.sources.patria_planilha_fundamentos import (
    FundamentosPlanilha,
    parse_resumo_credito,
    parse_resumo_tijolo,
)
from iip.sources.patria_planilha_fundamentos_harvester import (
    FetchedPlanilhaFundamentos,
    _planilha_category,
)


def _tijolo_workbook(*, start_col: int = 3, pl_format: str = "General"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resumo"
    labels_and_values = [
        ("Patrimônio Líquido", 2500.0),
        ("Valor de Mercado", 2000.0),
        ("Nº de Ativos", 12),
        ("Nº de Locatários", 30),
        ("ABL (m2)", 500000.0),
        ("WALE", 4.5),
        ("Vacância Física", 0.05),
        ("Vacância Financeira", 0.02),
        ("P/VP", 0.9),
        ("DY (Mercado)", 0.08),
        ("DY (Patrimonial)", 0.07),
    ]
    for i, (label, value) in enumerate(labels_and_values):
        col = start_col + i
        ws.cell(row=10, column=col, value=label)
        cell = ws.cell(row=11, column=col, value=value)
        if label == "Patrimônio Líquido":
            cell.number_format = pl_format
    return wb


def test_parse_resumo_tijolo_reads_values_below_labels():
    result = parse_resumo_tijolo(_tijolo_workbook(), "lvbi11")

    assert result is not None
    assert result.ticker == "LVBI11"
    assert result.wale_anos == 4.5
    assert result.vacancia_financeira == 0.02
    assert result.n_locatarios == 30
    assert result.occupancy_rate == pytest.approx(0.98)


def test_parse_resumo_tijolo_is_column_position_independent():
    result = parse_resumo_tijolo(_tijolo_workbook(start_col=2), "HGRU11")

    assert result is not None
    assert result.wale_anos == 4.5


@pytest.mark.parametrize(
    ("fmt", "expected"),
    [
        ('"R$ "#,##0.0" milhões"', 2_500_000_000.0),
        ('"R$ "#,##0.0" bilhões"', 2_500_000_000_000.0),
        ("General", 2500.0),
    ],
)
def test_parse_resumo_tijolo_normalizes_unit_from_number_format(fmt, expected):
    result = parse_resumo_tijolo(_tijolo_workbook(pl_format=fmt), "PVBI11")

    assert result is not None
    assert result.patrimonio_liquido == expected


def test_parse_resumo_tijolo_returns_none_without_resumo_sheet():
    wb = openpyxl.Workbook()
    wb.active.title = "Outra"

    assert parse_resumo_tijolo(wb, "HGRU11") is None


def test_parse_resumo_tijolo_returns_none_for_credit_fund_layout():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resumo"
    ws["B2"] = "Patrimônio Líquido"
    ws["B3"] = 100
    ws["D2"] = "% PL por classe de ativo"

    assert parse_resumo_tijolo(wb, "HGCR11") is None


def test_occupancy_rate_is_none_without_vacancia_financeira():
    fund = FundamentosPlanilha(
        ticker="X",
        competencia=None,
        patrimonio_liquido=None,
        valor_mercado=None,
        n_ativos=None,
        n_locatarios=None,
        abl_m2=None,
        wale_anos=None,
        vacancia_fisica=None,
        vacancia_financeira=None,
        p_vp=None,
        dy_mercado=None,
        dy_patrimonial=None,
    )

    assert fund.occupancy_rate is None


def test_planilha_category_matches_slug_and_label_styles():
    assert (
        _planilha_category(("lvbi11_relatorio", "lvbi11_planilha_de_fundamentos"))
        == "lvbi11_planilha_de_fundamentos"
    )
    assert (
        _planilha_category(("HGRU - Planilha de Fundamentos",))
        == "HGRU - Planilha de Fundamentos"
    )
    assert _planilha_category(("hgru_relatorio_gerencial",)) is None


# --- _enrich_fii_with_patria_fundamentos -------------------------------


def _fundamentos(**overrides) -> FundamentosPlanilha:
    base = {
        "ticker": "LVBI11",
        "competencia": None,
        "patrimonio_liquido": None,
        "valor_mercado": None,
        "n_ativos": None,
        "n_locatarios": None,
        "abl_m2": None,
        "wale_anos": 4.5678,
        "vacancia_fisica": 0.05,
        "vacancia_financeira": 0.02,
        "p_vp": None,
        "dy_mercado": None,
        "dy_patrimonial": None,
    }
    base.update(overrides)
    return FundamentosPlanilha(**base)


def _patch_harvester(monkeypatch, *, result=None, error=None):
    class _FakeHarvester:
        def fetch(self, ticker):
            if error is not None:
                raise error
            return result

    monkeypatch.setattr(
        harvester_module, "PatriaPlanilhaFundamentosHTTPHarvester", _FakeHarvester
    )


def test_enrich_is_silent_noop_for_non_patria_ticker(monkeypatch):
    _patch_harvester(monkeypatch, error=AssertionError("must not be called"))
    financials = {"occupancy_rate": 0.85}

    out, fetched, warnings = fetch_template._enrich_fii_with_patria_fundamentos(
        financials, "BTLG11"
    )

    assert out is financials
    assert fetched == []
    assert warnings == []


def test_enrich_fills_occupancy_and_wale(monkeypatch):
    _patch_harvester(
        monkeypatch,
        result=FetchedPlanilhaFundamentos(
            ticker="LVBI11", category="c", document=None, fundamentos=_fundamentos()
        ),
    )
    financials = {"occupancy_rate": 0.85, "avg_lease_term_years": 5, "other": 1}

    out, fetched, warnings = fetch_template._enrich_fii_with_patria_fundamentos(
        financials, "LVBI11"
    )

    assert out["occupancy_rate"] == pytest.approx(0.98)
    assert out["avg_lease_term_years"] == 4.57
    assert out["other"] == 1
    assert fetched == ["occupancy_rate", "avg_lease_term_years"]
    assert warnings == []
    assert financials["occupancy_rate"] == 0.85  # input not mutated


def test_enrich_warns_when_layout_is_unrecognized(monkeypatch):
    _patch_harvester(
        monkeypatch,
        result=FetchedPlanilhaFundamentos(
            ticker="HGCR11", category="c", document=None, fundamentos=None
        ),
    )
    financials = {"occupancy_rate": 0.85}

    out, fetched, warnings = fetch_template._enrich_fii_with_patria_fundamentos(
        financials, "HGCR11"
    )

    assert out == financials
    assert fetched == []
    assert len(warnings) == 1


def test_enrich_never_raises_on_fetch_failure(monkeypatch):
    _patch_harvester(monkeypatch, error=OSError("rede caiu"))

    out, fetched, warnings = fetch_template._enrich_fii_with_patria_fundamentos(
        {"occupancy_rate": 0.85}, "LVBI11"
    )

    assert out == {"occupancy_rate": 0.85}
    assert fetched == []
    assert "rede caiu" in warnings[0]


# --- credit-fund layout (HGCR11 / PCIP11) -------------------------------


def _credito_workbook(*, c0: int = 2, carteira_prazo=3.67):
    """Mimics the real HGCR11 (c0=2) / PCIP11 (c0=3) "Resumo" sheets: PL and
    valor de mercado carry their value to the RIGHT of the label, the
    per-cota indicators sit one row BELOW their header, and the
    composition table is keyed by a "% PL" header with its row labels one
    column to the left."""

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resumo"
    ws.cell(row=4, column=c0, value="Fundo Crédito")
    ws.cell(row=4, column=c0 + 1, value=datetime.datetime(2026, 7, 31))

    ws.cell(row=7, column=c0, value="Patrimônio líquido")
    pl = ws.cell(row=7, column=c0 + 1, value=1500.0)
    pl.number_format = r'"R$"\ #,##0.0" milhões"'
    ws.cell(row=7, column=c0 + 2, value=100.0)
    ws.cell(row=8, column=c0, value="Valor de Mercado")
    vm = ws.cell(row=8, column=c0 + 1, value=1400.0)
    vm.number_format = r'"R$"\ #,##0.0" milhões"'
    ws.cell(row=8, column=c0 + 2, value=94.0)

    ws.cell(row=7, column=c0 + 6, value="Rendimento \npor Cota")
    ws.cell(row=8, column=c0 + 6, value=1.0)
    ws.cell(row=7, column=c0 + 7, value="Reserva \nAcumulada")
    ws.cell(row=8, column=c0 + 7, value=0.7)
    ws.cell(row=7, column=c0 + 8, value="Número de \nCotistas")
    ws.cell(row=8, column=c0 + 8, value=100000)

    ws.cell(row=11, column=c0, value="Tabela de Sensibilidade")
    headers = [
        "% PL",
        "Yield \nNominal",
        "IPCA\n Ref.",
        "Yield \n(IPCA +)",
        "Prazo Médio (Anos)",
        "Spread",
    ]
    for i, h in enumerate(headers):
        ws.cell(row=12, column=c0 + 6 + i, value=h)
    table = [
        ("CRI + Op. Estrut", [0.89, 0.16, 0.05, 0.104, 3.6, 0.02]),
        ("FII", [0.08, 0.16, 0.05, 0.102, 4, 0.018]),
        ("Carteira", [0.97, 0.16, 0, 0.1042, carteira_prazo, 0.0199]),
        ("Caixa", [0.03, 0.11, 0.04, 0.065, "-", "-"]),
        ("Compromissada", ["-"] * 6),
    ]
    for r, (label, values) in enumerate(table, start=13):
        ws.cell(row=r, column=c0 + 5, value=label)
        for i, v in enumerate(values):
            ws.cell(row=r, column=c0 + 6 + i, value=v)
    return wb


@pytest.mark.parametrize("c0", [2, 3])
def test_parse_resumo_credito_reads_real_layout_at_any_column_offset(c0):
    result = parse_resumo_credito(_credito_workbook(c0=c0), "hgcr11")

    assert result is not None
    assert result.ticker == "HGCR11"
    assert result.competencia == datetime.date(2026, 7, 31)
    assert result.patrimonio_liquido == 1_500_000_000.0
    assert result.valor_mercado == 1_400_000_000.0
    assert result.vp_cota == 100.0
    assert result.preco_cota == 94.0
    assert result.rendimento_cota == 1.0
    assert result.reserva_acumulada_cota == 0.7
    assert result.n_cotistas == 100000
    assert result.pct_pl_cri == 0.89
    assert result.pct_pl_fii == 0.08
    assert result.pct_pl_caixa == 0.03
    assert result.yield_ipca_carteira == 0.1042
    assert result.prazo_medio_carteira_anos == 3.67
    assert result.spread_carteira == 0.0199


def test_parse_resumo_credito_treats_dash_as_none():
    result = parse_resumo_credito(_credito_workbook(carteira_prazo="-"), "HGCR11")

    assert result is not None
    assert result.prazo_medio_carteira_anos is None


def test_credito_and_tijolo_parsers_reject_each_others_layout():
    assert parse_resumo_credito(_tijolo_workbook(), "LVBI11") is None
    assert parse_resumo_tijolo(_credito_workbook(), "HGCR11") is None


def test_parse_resumo_credito_returns_none_without_resumo_sheet():
    wb = openpyxl.Workbook()
    wb.active.title = "Outra"

    assert parse_resumo_credito(wb, "HGCR11") is None


def test_reserves_to_npa_divides_reserve_by_vp_and_guards_missing_inputs():
    result = parse_resumo_credito(_credito_workbook(), "HGCR11")

    assert result.reserves_to_npa == pytest.approx(0.007)
    assert dataclasses.replace(result, vp_cota=None).reserves_to_npa is None
    assert dataclasses.replace(result, vp_cota=0).reserves_to_npa is None


def test_enrich_credit_fund_fills_only_reserves_to_npa(monkeypatch):
    credito = parse_resumo_credito(_credito_workbook(), "HGCR11")
    _patch_harvester(
        monkeypatch,
        result=FetchedPlanilhaFundamentos(
            ticker="HGCR11",
            category="c",
            document=None,
            fundamentos=None,
            credito=credito,
        ),
    )
    financials = {
        "occupancy_rate": 0.85,
        "avg_lease_term_years": 5,
        "reserves_to_npa": 0.05,
    }

    out, fetched, warnings = fetch_template._enrich_fii_with_patria_fundamentos(
        financials, "HGCR11"
    )

    assert out["reserves_to_npa"] == pytest.approx(0.007)
    # CRI duration is not a lease term: those defaults must stay untouched.
    assert out["occupancy_rate"] == 0.85
    assert out["avg_lease_term_years"] == 5
    assert fetched == ["reserves_to_npa"]
    assert warnings == []
    assert financials["reserves_to_npa"] == 0.05  # input not mutated
